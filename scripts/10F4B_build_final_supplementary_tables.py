#!/usr/bin/env python3

from pathlib import Path
import csv
import hashlib
import re
import sys
from datetime import datetime

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import (
        Font, PatternFill, Border, Side,
        Alignment
    )
    from openpyxl.utils import get_column_letter
except Exception as e:
    raise SystemExit(
        "ERROR: openpyxl is required.\n"
        "Install with:\n"
        "python3 -m pip install --user openpyxl\n"
        f"{e}"
    )

print("=" * 88)
print("STEP10F4B — FINAL SUPPLEMENTARY TABLE WORKBOOK CONSTRUCTION")
print("Build from Step10F4A2 frozen sources only")
print("No biological analysis rerun")
print("No statistic recomputed")
print("=" * 88)

ROOT = Path(".")

LOCKDIR = Path(
    "10_publication_figures/"
    "Step10F_final_consistency/"
    "Step10F4_Supplementary_Tables/"
    "Step10F4A2_final_source_lock"
)

LOCK_MANIFEST = (
    LOCKDIR /
    "10F4A2_final_component_source_lock.tsv"
)

OUTDIR = Path(
    "10_publication_figures/"
    "Step10F_final_consistency/"
    "Step10F4_Supplementary_Tables/"
    "Step10F4B_final_workbooks"
)

OUTDIR.mkdir(
    parents=True,
    exist_ok=True
)

if not LOCK_MANIFEST.exists():
    raise SystemExit(
        f"ERROR: source-lock manifest not found:\n{LOCK_MANIFEST}"
    )


# =====================================================================
# Frozen final workbook definitions
# =====================================================================

WORKBOOK_INFO = {

    "S1": {
        "filename": "Supplementary_Table_S1.xlsx",
        "title":
            "Supplementary Table S1. Paired multi-tissue RNA-seq and ATAC-seq sample metadata.",
        "description":
            "Paired sample metadata linking tissue, animal identity, RNA-seq sample and ATAC-seq sample.",
        "notes": [
            "RNA-seq and ATAC-seq were obtained from the same biological animals P348 and P350 across eight tissues.",
            "SRA accessions are retained for traceability; figure sample labels use the corresponding tissue–animal identities.",
            "No biological values were modified during workbook construction."
        ],
        "include": [
            "Sample_metadata"
        ]
    },

    "S2": {
        "filename": "Supplementary_Table_S2.xlsx",
        "title":
            "Supplementary Table S2. RNA-seq and ATAC-seq quality control, multi-tissue effects and tissue specificity.",
        "description":
            "Frozen summary outputs from sample filtering, DESeq2 multi-tissue LRT, Tau tissue specificity and SPM analyses.",
        "notes": [
            "The DESeq2 multi-tissue test used the full model ~ Animal + Tissue and reduced model ~ Animal.",
            "FDR < 0.01 was used as the downstream multi-tissue significance threshold.",
            "Tau and SPM were calculated from linear tissue-mean TMM.TPM values; SPM is treated as a continuous tissue-preference measure.",
            "High-confidence categorical tissue-specific membership is taken from the frozen Step03 output."
        ],
        "include": [
            "RNA_filtering",
            "ATAC_filtering",
            "Multi_tissue_LRT",
            "Tissue_specificity",
            "RNA_tissue_counts",
            "ATAC_tissue_counts",
            "RNA_SPM_summary",
            "ATAC_SPM_summary",
            "RNA_threshold_sensitivity",
            "ATAC_threshold_sensitivity"
        ]
    },

    "S3": {
        "filename": "Supplementary_Table_S3.xlsx",
        "title":
            "Supplementary Table S3. Strongly coupled tissue-specific peak–gene associations.",
        "description":
            "Frozen multi-evidence master table for the 702 strongly coupled tissue-specific peak–gene candidates.",
        "notes": [
            "Strong coupling required rho_P348 >= 0.5, rho_P350 >= 0.5 and rho_mean >= 0.7 among concordant high-confidence tissue-specific peak–gene associations.",
            "Peak–gene links are genomic associations based on the frozen UROPA assignment framework and should not be interpreted as experimentally validated causal interactions.",
            "The 702 rows are the frozen publication candidate set."
        ],
        "include": [
            "Strong_peak_gene_pairs"
        ]
    },

    "S4": {
        "filename": "Supplementary_Table_S4.xlsx",
        "title":
            "Supplementary Table S4. Genomic architecture and promoter-first annotation of strong peak–gene associations.",
        "description":
            "Frozen structural and promoter/TSS architecture results for the 702 strong peak–gene associations.",
        "notes": [
            "Promoters were defined relative to the representative-transcript TSS using a strand-aware -2000/+500 bp window.",
            "Promoter-first annotation priority was Promoter > 5UTR > 3UTR > Exon > Intron > Distal.",
            "Structural enrichment and promoter enrichment results reproduce the frozen Step06/Step07 outputs.",
            "Pair-specific genomic classes describe the associated peak–gene pair and do not establish regulatory causality."
        ],
        "include": [
            "Strong_structure",
            "UROPA_location_summary",
            "Structural_class_summary",
            "Structural_enrichment",
            "Promoter_first_702",
            "Promoter_class_summary",
            "TSS_distance_summary",
            "Promoter_summary",
            "Global_enrichment",
            "Pair_specific_enrichment"
        ]
    },

    "S5": {
        "filename": "Supplementary_Table_S5.xlsx",
        "title":
            "Supplementary Table S5. Matched-background motif analysis and motif reproducibility evidence.",
        "description":
            "Frozen matched-background QC, AME, STREME and Tomtom motif-analysis outputs.",
        "notes": [
            "Matched backgrounds were selected within tissue using pair-specific genomic class, GC content and sequence length.",
            "AME formal significance criterion was E < 0.05; no motif met this formal criterion.",
            "Default STREME reported 14 motifs; nine had evaluable holdout E-values and 0/9 met E < 0.05. Five Cerebellum motifs were training-score-only and therefore not formally evaluable by holdout E-value.",
            "No-holdout STREME and cross-run Tomtom similarity are supportive sensitivity/reproducibility evidence, not independent replication.",
            "Raw STREME text files remain source-locked and are listed in SOURCE_INDEX but are not duplicated as worksheet text."
        ],
        "include": [
            "Matched_background_QC",
            "AME_QC",
            "AME_all_results",
            "STREME_default_summary",
            "STREME_noholdout_summary",
            "Tomtom_repro_summary",
            "Tomtom_Cerebellum",
            "Tomtom_Liver",
            "Tomtom_Muscle",
            "Tomtom_Spleen"
        ]
    },

    "S6": {
        "filename": "Supplementary_Table_S6.xlsx",
        "title":
            "Supplementary Table S6. High-confidence motif–peak–gene sequence-match evidence.",
        "description":
            "Frozen FIMO sequence-match evidence and supporting motif/TF integration.",
        "notes": [
            "The PRIMARY sheet contains the 612 FIMO q <= 0.05 motif–peak–gene sequence-match links.",
            "FIMO evidence indicates sequence matching and must not be interpreted as direct TF occupancy or experimentally validated binding.",
            "The High_RNA_integrated sheet is supporting TF-expression integration and is not the primary FIMO evidence table.",
            "Motif-family assignment does not by itself establish binding by a specific TF."
        ],
        "include": [
            "FIMO_q005_links",
            "Motif_target_summary",
            "Tissue_network_summary",
            "High_RNA_integrated",
            "FIMO_QC"
        ]
    },

    "S7": {
        "filename": "Supplementary_Table_S7.xlsx",
        "title":
            "Supplementary Table S7. Candidate transcription-factor RNA-support master.",
        "description":
            "Frozen publication candidate TF master integrating motif-family and RNA-expression support.",
        "notes": [
            "The table contains 48 frozen tissue–TF candidate records.",
            "RNA support and motif-family evidence are used to prioritize candidate TFs but do not demonstrate direct DNA binding.",
            "Candidate TF terminology is retained throughout."
        ],
        "include": [
            "Candidate_TF_master"
        ]
    },

    "S8": {
        "filename": "Supplementary_Table_S8.xlsx",
        "title":
            "Supplementary Table S8. Candidate regulatory modules and representative network support.",
        "description":
            "Frozen candidate-module definitions, peak–gene candidates, gene rankings and representative network files.",
        "notes": [
            "Candidate regulatory modules summarize convergent multi-evidence patterns and should not be interpreted as validated direct regulatory interactions.",
            "Within-tissue motif overlaps describe shared candidate peaks/genes and do not establish TF cooperation.",
            "The representative network is a visualization-oriented subset of the frozen candidate network."
        ],
        "include": [
            "Module_summary",
            "Module_definitions",
            "Module_peak_gene_pairs",
            "Gene_ranking",
            "Top10_representatives",
            "Top5_primary_supporting",
            "Within_tissue_overlap",
            "Representative_edges",
            "Representative_nodes"
        ]
    }
}


# =====================================================================
# Helpers
# =====================================================================

def file_md5(path):

    h = hashlib.md5()

    with open(path, "rb") as f:
        for block in iter(
            lambda: f.read(1024 * 1024),
            b""
        ):
            h.update(block)

    return h.hexdigest()


def safe_sheet_name(name):

    name = re.sub(
        r'[\[\]\:\*\?\/\\]',
        "_",
        name
    )

    return name[:31]


TEXT_HEADER_KEYWORDS = [
    "id",
    "name",
    "symbol",
    "tissue",
    "animal",
    "sample",
    "chr",
    "chrom",
    "strand",
    "feature",
    "class",
    "location",
    "role",
    "source",
    "label",
    "consensus",
    "sequence",
    "motif",
    "tf",
    "gene",
    "peak",
    "bound",
    "module"
]


def should_keep_text(header):

    h = header.lower()

    return any(
        x in h
        for x in TEXT_HEADER_KEYWORDS
    )


INT_RE = re.compile(
    r"^[+-]?\d+$"
)

FLOAT_RE = re.compile(
    r"^[+-]?(?:"
    r"(?:\d+\.\d*)|"
    r"(?:\d*\.\d+)|"
    r"(?:\d+)"
    r")(?:[eE][+-]?\d+)?$"
)


def convert_value(value, header):

    if value is None:
        return None

    s = str(value)

    if s == "":
        return None

    # Preserve explicit missing / special tokens.
    if s.lower() in {
        "na",
        "nan",
        "inf",
        "-inf",
        "null",
        "none"
    }:
        return s

    if should_keep_text(header):
        return s

    if INT_RE.match(s):

        try:
            return int(s)
        except Exception:
            return s

    if FLOAT_RE.match(s):

        try:
            return float(s)
        except Exception:
            return s

    return s


def read_delimited(path):

    delimiter = (
        "\t"
        if path.suffix.lower() == ".tsv"
        else ","
    )

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace",
        newline=""
    ) as f:

        reader = csv.reader(
            f,
            delimiter=delimiter
        )

        rows = []

        for row in reader:

            if not row:
                continue

            if (
                len(row) == 1
                and row[0].startswith("#")
            ):
                continue

            rows.append(row)

    if not rows:
        return [], []

    headers = rows[0]

    converted = []

    for row in rows[1:]:

        # Pad short rows if necessary.
        if len(row) < len(headers):
            row = row + [""] * (
                len(headers) - len(row)
            )

        # Do not silently keep unexpected overflow.
        if len(row) > len(headers):
            raise ValueError(
                f"Column mismatch in {path}: "
                f"header={len(headers)}, row={len(row)}"
            )

        converted.append(
            [
                convert_value(
                    v,
                    headers[i]
                )
                for i, v in enumerate(row)
            ]
        )

    return headers, converted


def number_format_for_header(header):

    h = header.lower()

    if any(
        x in h
        for x in [
            "p_value",
            "pvalue",
            "p.value",
            "adj_p",
            "fdr",
            "q_value",
            "qvalue",
            "best_p",
            "best_q",
            "e_value",
            "evalue"
        ]
    ):
        return "0.00E+00"

    if any(
        x in h
        for x in [
            "rho",
            "tau",
            "spm",
            "percent",
            "fraction",
            "ratio",
            "fold",
            "odds",
            "score",
            "identity",
            "similarity"
        ]
    ):
        return "0.000"

    if "tpm" in h:
        return "0.000"

    if any(
        x in h
        for x in [
            "distance",
            "start",
            "end"
        ]
    ):
        return "0"

    return None


# =====================================================================
# Styles
# =====================================================================

HEADER_FILL = PatternFill(
    "solid",
    fgColor="1F4E78"
)

SECONDARY_FILL = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

README_FILL = PatternFill(
    "solid",
    fgColor="17365D"
)

WHITE_FONT = Font(
    color="FFFFFF",
    bold=True
)

HEADER_FONT = Font(
    color="FFFFFF",
    bold=True,
    size=10
)

THIN_GREY = Side(
    style="thin",
    color="D9E1F2"
)

BOTTOM_BORDER = Border(
    bottom=THIN_GREY
)


def style_data_sheet(ws):

    max_row = ws.max_row
    max_col = ws.max_column

    if max_row < 1 or max_col < 1:
        return

    # Header.
    for cell in ws[1]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        cell.border = BOTTOM_BORDER

    ws.row_dimensions[1].height = 36

    # Freeze header and useful identifier columns.
    if max_col >= 20:
        ws.freeze_panes = "D2"
    elif max_col >= 8:
        ws.freeze_panes = "B2"
    else:
        ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        f"A1:{get_column_letter(max_col)}{max_row}"
    )

    # Column widths based primarily on headers and a limited sample.
    sample_n = min(
        max_row,
        60
    )

    for col_idx in range(
        1,
        max_col + 1
    ):

        header = str(
            ws.cell(
                1,
                col_idx
            ).value or ""
        )

        lengths = [
            len(header)
        ]

        for r in range(
            2,
            sample_n + 1
        ):

            v = ws.cell(
                r,
                col_idx
            ).value

            if v is not None:
                lengths.append(
                    len(str(v))
                )

        width = max(lengths) + 2

        # Long text columns.
        h = header.lower()

        if any(
            x in h
            for x in [
                "description",
                "candidate_tfs",
                "evidence_sources",
                "consensus"
            ]
        ):
            width = max(
                width,
                24
            )

        width = min(
            max(width, 10),
            34
        )

        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

        fmt = number_format_for_header(
            header
        )

        if fmt:

            for r in range(
                2,
                max_row + 1
            ):

                cell = ws.cell(
                    r,
                    col_idx
                )

                if isinstance(
                    cell.value,
                    (int, float)
                ):
                    cell.number_format = fmt

    ws.sheet_view.showGridLines = False


def add_readme_sheet(
    wb,
    table_id,
    info,
    source_rows
):

    ws = wb.create_sheet(
        "README"
    )

    ws.merge_cells(
        "A1:H1"
    )

    ws["A1"] = info["title"]
    ws["A1"].fill = README_FILL
    ws["A1"].font = Font(
        color="FFFFFF",
        bold=True,
        size=15
    )
    ws["A1"].alignment = Alignment(
        vertical="center",
        wrap_text=True
    )

    ws.row_dimensions[1].height = 36

    metadata = [
        ("Table", table_id),
        (
            "Construction status",
            "FINAL workbook built from Step10F4A2 frozen sources"
        ),
        (
            "Biological results modified",
            "0"
        ),
        (
            "Description",
            info["description"]
        ),
        (
            "Source lock",
            str(
                LOCK_MANIFEST
            )
        ),
        (
            "Construction date",
            datetime.now().strftime(
                "%Y-%m-%d"
            )
        )
    ]

    row = 3

    for key, value in metadata:

        ws.cell(
            row,
            1,
            key
        )

        ws.cell(
            row,
            2,
            value
        )

        ws.cell(
            row,
            1
        ).font = Font(
            bold=True
        )

        ws.cell(
            row,
            1
        ).fill = SECONDARY_FILL

        ws.cell(
            row,
            2
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

        row += 1

    row += 1

    ws.cell(
        row,
        1,
        "Interpretive notes"
    )

    ws.cell(
        row,
        1
    ).font = Font(
        bold=True,
        color="FFFFFF"
    )

    ws.cell(
        row,
        1
    ).fill = HEADER_FILL

    row += 1

    for i, note in enumerate(
        info["notes"],
        start=1
    ):

        ws.cell(
            row,
            1,
            i
        )

        ws.cell(
            row,
            2,
            note
        )

        ws.cell(
            row,
            2
        ).alignment = Alignment(
            wrap_text=True,
            vertical="top"
        )

        row += 1

    row += 1

    ws.cell(
        row,
        1,
        "Workbook sheets"
    )

    ws.cell(
        row,
        1
    ).font = Font(
        bold=True,
        color="FFFFFF"
    )

    ws.cell(
        row,
        1
    ).fill = HEADER_FILL

    row += 1

    ws.append(
        [
            "Sheet",
            "Role",
            "Description"
        ]
    )

    hdr = row

    for c in range(
        1,
        4
    ):
        ws.cell(
            hdr,
            c
        ).font = WHITE_FONT

        ws.cell(
            hdr,
            c
        ).fill = HEADER_FILL

    by_sheet = {
        x["Sheet"]: x
        for x in source_rows
    }

    for sheet in info["include"]:

        x = by_sheet[sheet]

        ws.append(
            [
                sheet,
                x["Role"],
                x["Description"]
            ]
        )

    ws.column_dimensions["A"].width = 31
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 95

    for row_cells in ws.iter_rows():

        for cell in row_cells:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    return ws


def add_source_index(
    wb,
    source_rows,
    included
):

    ws = wb.create_sheet(
        "SOURCE_INDEX"
    )

    headers = [
        "Table",
        "Sheet",
        "Role",
        "Source_path",
        "Expected_rows",
        "Observed_rows",
        "Locked_MD5",
        "Included_in_workbook",
        "Description"
    ]

    ws.append(headers)

    for x in source_rows:

        ws.append(
            [
                x["Table"],
                x["Sheet"],
                x["Role"],
                x["Path"],
                x["Expected_rows"],
                x["Observed_rows"],
                x["MD5"],
                (
                    "YES"
                    if x["Sheet"] in included
                    else "SOURCE_LOCK_ONLY"
                ),
                x["Description"]
            ]
        )

    style_data_sheet(
        ws
    )

    ws.column_dimensions["D"].width = 75
    ws.column_dimensions["G"].width = 34
    ws.column_dimensions["I"].width = 80

    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    return ws


# =====================================================================
# Read and verify source lock
# =====================================================================

with open(
    LOCK_MANIFEST,
    "r",
    encoding="utf-8"
) as f:

    locked_rows = list(
        csv.DictReader(
            f,
            delimiter="\t"
        )
    )

if len(locked_rows) != 55:

    raise SystemExit(
        "ERROR: expected 55 source-lock rows, "
        f"found {len(locked_rows)}."
    )


print()
print("=" * 88)
print("VERIFYING LOCKED SOURCE MD5")
print("=" * 88)

md5_pass = 0

for x in locked_rows:

    p = ROOT / x["Path"]

    if not p.exists():

        raise SystemExit(
            f"ERROR: locked source missing:\n{p}"
        )

    current = file_md5(p)

    if current != x["MD5"]:

        raise SystemExit(
            "ERROR: source changed after Step10F4A2 lock:\n"
            f"{p}\n"
            f"Locked : {x['MD5']}\n"
            f"Current: {current}"
        )

    md5_pass += 1

print(
    f"Locked source MD5 verification: "
    f"{md5_pass}/55 PASS"
)


# =====================================================================
# Build workbooks
# =====================================================================

build_manifest = []

sheet_qc = []

for table_id in [
    "S1",
    "S2",
    "S3",
    "S4",
    "S5",
    "S6",
    "S7",
    "S8"
]:

    info = WORKBOOK_INFO[
        table_id
    ]

    source_rows = [
        x for x in locked_rows
        if x["Table"] == table_id
    ]

    source_by_sheet = {
        x["Sheet"]: x
        for x in source_rows
    }

    # Ensure final planned sheets exist in source lock.
    missing = [
        x for x in info["include"]
        if x not in source_by_sheet
    ]

    if missing:

        raise SystemExit(
            f"ERROR: {table_id} planned sheets "
            f"not present in source lock:\n"
            + "\n".join(missing)
        )

    print()
    print(
        f"Building {table_id}: "
        f"{info['filename']}"
    )

    wb = Workbook()

    default = wb.active
    wb.remove(default)

    wb.properties.title = (
        info["title"]
    )

    wb.properties.subject = (
        info["description"]
    )

    wb.properties.creator = (
        "Step10F4B frozen-source workbook construction"
    )

    # README first.
    add_readme_sheet(
        wb,
        table_id,
        info,
        source_rows
    )

    # Source index second.
    add_source_index(
        wb,
        source_rows,
        set(
            info["include"]
        )
    )

    # Data sheets.
    for sheet_name in info["include"]:

        source = source_by_sheet[
            sheet_name
        ]

        path = ROOT / source["Path"]

        if path.suffix.lower() not in {
            ".tsv",
            ".csv"
        }:

            raise SystemExit(
                f"ERROR: non-tabular source selected "
                f"for workbook sheet: {path}"
            )

        headers, rows = read_delimited(
            path
        )

        ws = wb.create_sheet(
            safe_sheet_name(
                sheet_name
            )
        )

        ws.append(
            headers
        )

        for row in rows:
            ws.append(row)

        style_data_sheet(
            ws
        )

        expected = source[
            "Expected_rows"
        ]

        if expected == "EXISTENCE_ONLY":
            expected_n = None
        else:
            expected_n = int(
                expected
            )

        observed_n = len(rows)

        status = (
            "PASS"
            if (
                expected_n is None
                or observed_n == expected_n
            )
            else "FAIL"
        )

        sheet_qc.append(
            {
                "Table":
                    table_id,
                "Sheet":
                    sheet_name,
                "Source":
                    source["Path"],
                "Expected_rows":
                    (
                        expected_n
                        if expected_n is not None
                        else "NA"
                    ),
                "Workbook_data_rows":
                    observed_n,
                "Columns":
                    len(headers),
                "Status":
                    status
            }
        )

        if status != "PASS":
            raise SystemExit(
                f"ERROR: row-count mismatch in "
                f"{table_id}/{sheet_name}"
            )

    outfile = (
        OUTDIR /
        info["filename"]
    )

    wb.save(
        outfile
    )

    # Re-open for structural verification.
    check_wb = load_workbook(
        outfile,
        read_only=True,
        data_only=False
    )

    expected_sheet_names = [
        "README",
        "SOURCE_INDEX"
    ] + [
        safe_sheet_name(x)
        for x in info["include"]
    ]

    if check_wb.sheetnames != expected_sheet_names:

        raise SystemExit(
            f"ERROR: workbook sheet-order mismatch: {outfile}\n"
            f"Expected: {expected_sheet_names}\n"
            f"Found:    {check_wb.sheetnames}"
        )

    sheet_count = len(
        check_wb.sheetnames
    )

    check_wb.close()

    build_manifest.append(
        {
            "Table":
                table_id,
            "Workbook":
                str(outfile),
            "Workbook_MD5":
                file_md5(outfile),
            "Size_bytes":
                outfile.stat().st_size,
            "Sheet_count":
                sheet_count,
            "Status":
                "PASS"
        }
    )

    print(
        f"  PASS: {sheet_count} sheets"
    )


# =====================================================================
# Primary frozen-count QC
# =====================================================================

primary_expectations = {
    ("S1", "Sample_metadata"): 16,

    ("S3", "Strong_peak_gene_pairs"): 702,

    ("S4", "Strong_structure"): 702,
    ("S4", "Promoter_first_702"): 702,

    ("S6", "FIMO_q005_links"): 612,

    ("S7", "Candidate_TF_master"): 48,

    ("S8", "Module_summary"): 11,
    ("S8", "Module_peak_gene_pairs"): 787,
    ("S8", "Gene_ranking"): 596,
    ("S8", "Top10_representatives"): 89,
    ("S8", "Top5_primary_supporting"): 30,
    ("S8", "Representative_edges"): 119,
    ("S8", "Representative_nodes"): 73
}

primary_qc = []

for (
    table_id,
    sheet
), expected in primary_expectations.items():

    hit = [
        x for x in sheet_qc
        if (
            x["Table"] == table_id
            and
            x["Sheet"] == sheet
        )
    ]

    if len(hit) != 1:

        observed = "NA"
        status = "FAIL"

    else:

        observed = hit[0][
            "Workbook_data_rows"
        ]

        status = (
            "PASS"
            if observed == expected
            else "FAIL"
        )

    primary_qc.append(
        {
            "Metric":
                f"{table_id}_{sheet}",
            "Observed":
                observed,
            "Expected":
                expected,
            "Status":
                status
        }
    )

    if status != "PASS":
        raise SystemExit(
            f"ERROR: primary QC failed: "
            f"{table_id}/{sheet}"
        )


# =====================================================================
# Write manifests
# =====================================================================

manifest_file = (
    OUTDIR /
    "Supplementary_Tables_manifest.tsv"
)

with open(
    manifest_file,
    "w",
    newline=""
) as f:

    fields = [
        "Table",
        "Workbook",
        "Workbook_MD5",
        "Size_bytes",
        "Sheet_count",
        "Status"
    ]

    writer = csv.DictWriter(
        f,
        fieldnames=fields,
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        build_manifest
    )


sheet_qc_file = (
    OUTDIR /
    "Step10F4B_sheet_QC.tsv"
)

with open(
    sheet_qc_file,
    "w",
    newline=""
) as f:

    fields = [
        "Table",
        "Sheet",
        "Source",
        "Expected_rows",
        "Workbook_data_rows",
        "Columns",
        "Status"
    ]

    writer = csv.DictWriter(
        f,
        fieldnames=fields,
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        sheet_qc
    )


primary_qc_file = (
    OUTDIR /
    "Step10F4B_primary_count_QC.tsv"
)

with open(
    primary_qc_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Metric",
            "Observed",
            "Expected",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        primary_qc
    )


md5_file = (
    OUTDIR /
    "Supplementary_Tables_MD5.tsv"
)

with open(
    md5_file,
    "w"
) as f:

    f.write(
        "Table\tWorkbook\tMD5\n"
    )

    for x in build_manifest:

        f.write(
            f"{x['Table']}\t"
            f"{x['Workbook']}\t"
            f"{x['Workbook_MD5']}\n"
        )


overall_file = (
    OUTDIR /
    "Step10F4B_overall_status.tsv"
)

with open(
    overall_file,
    "w"
) as f:

    f.write(
        "Metric\tValue\n"
    )

    f.write(
        "Locked_source_files_verified\t55\n"
    )

    f.write(
        "Locked_source_MD5_pass\t55\n"
    )

    f.write(
        "Final_workbooks_generated\t8\n"
    )

    f.write(
        "Primary_count_QC\tPASS\n"
    )

    f.write(
        "Step10F4B_final_workbooks\tPASS\n"
    )

    f.write(
        "Biological_results_modified\t0\n"
    )


# =====================================================================
# Console report
# =====================================================================

print()
print("=" * 88)
print("FINAL WORKBOOK SUMMARY")
print("=" * 88)

for x in build_manifest:

    mb = (
        x["Size_bytes"] /
        1024 /
        1024
    )

    print(
        f"{x['Table']}  "
        f"sheets={x['Sheet_count']:<3}  "
        f"size={mb:7.2f} MB  "
        f"PASS"
    )


print()
print("=" * 88)
print("PRIMARY DATA COUNT QC")
print("=" * 88)

for x in primary_qc:

    print(
        f"{x['Metric']:<43} "
        f"{str(x['Observed']):>6} / "
        f"{str(x['Expected']):<6} "
        f"{x['Status']}"
    )


print()
print("=" * 88)
print("STEP10F4B STATUS: PASS")
print("Final Supplementary Tables S1-S8 generated.")
print("All 55 locked source MD5 values verified before construction.")
print("Biological results modified: 0")
print("=" * 88)

print("\nOutput directory:")
print(OUTDIR)

print("\nFiles:")
for x in build_manifest:
    print(x["Workbook"])

print(manifest_file)
print(sheet_qc_file)
print(primary_qc_file)
print(md5_file)
print(overall_file)
