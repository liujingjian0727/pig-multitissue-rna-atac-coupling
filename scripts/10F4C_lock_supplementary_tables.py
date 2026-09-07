#!/usr/bin/env python3

from pathlib import Path
import csv
import hashlib
import sys
import re

try:
    from openpyxl import load_workbook
except Exception as e:
    raise SystemExit(
        "ERROR: openpyxl is required.\n"
        f"{e}"
    )

print("=" * 92)
print("STEP10F4C — SUPPLEMENTARY TABLES S1–S8 CONSISTENCY LOCK")
print("Technical / source-traceability / semantic consistency only")
print("No workbook reconstruction")
print("No statistic recomputed")
print("No biological result modified")
print("=" * 92)

# =====================================================================
# Paths
# =====================================================================

ROOT = Path(".")

BASE = Path(
    "10_publication_figures/"
    "Step10F_final_consistency/"
    "Step10F4_Supplementary_Tables"
)

LOCK_A2 = (
    BASE /
    "Step10F4A2_final_source_lock"
)

BUILD_B = (
    BASE /
    "Step10F4B_final_workbooks"
)

OUTDIR = (
    BASE /
    "Step10F4C_consistency_lock"
)

OUTDIR.mkdir(
    parents=True,
    exist_ok=True
)

SOURCE_LOCK = (
    LOCK_A2 /
    "10F4A2_final_component_source_lock.tsv"
)

BUILD_MANIFEST = (
    BUILD_B /
    "Supplementary_Tables_manifest.tsv"
)

BUILD_SHEET_QC = (
    BUILD_B /
    "Step10F4B_sheet_QC.tsv"
)

BUILD_PRIMARY_QC = (
    BUILD_B /
    "Step10F4B_primary_count_QC.tsv"
)

BUILD_MD5 = (
    BUILD_B /
    "Supplementary_Tables_MD5.tsv"
)

for f in [
    SOURCE_LOCK,
    BUILD_MANIFEST,
    BUILD_SHEET_QC,
    BUILD_PRIMARY_QC,
    BUILD_MD5
]:
    if not f.exists():
        raise SystemExit(
            f"ERROR: required frozen manifest missing:\n{f}"
        )


# =====================================================================
# Expected workbooks
# =====================================================================

EXPECTED = {

    "S1": {
        "file":
            "Supplementary_Table_S1.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Sample_metadata"
        ],

        "source_n": 1
    },

    "S2": {
        "file":
            "Supplementary_Table_S2.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "RNA_filtering",
            "ATAC_filtering",
            "Multi_tissue_LRT",
            "Tissue_specificity",
            "RNA_tissue_counts",
            "ATAC_tissue_counts",
            "RNA_SPM_summary",
            "ATAC_SPM_summary",
            "RNA_threshold_sensitivity",
            "ATAC_threshold_sensitivity"
        ],

        "source_n": 10
    },

    "S3": {
        "file":
            "Supplementary_Table_S3.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Strong_peak_gene_pairs"
        ],

        "source_n": 1
    },

    "S4": {
        "file":
            "Supplementary_Table_S4.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Strong_structure",
            "UROPA_location_summary",
            "Structural_class_summary",
            "Structural_enrichment",
            "Promoter_first_702",
            "Promoter_class_summary",
            "TSS_distance_summary",
            "Promoter_summary",
            "Global_enrichment",
            "Pair_specific_enrichment"
        ],

        "source_n": 10
    },

    "S5": {
        "file":
            "Supplementary_Table_S5.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Matched_background_QC",
            "AME_QC",
            "AME_all_results",
            "STREME_default_summary",
            "STREME_noholdout_summary",
            "Tomtom_repro_summary",
            "Tomtom_Cerebellum",
            "Tomtom_Liver",
            "Tomtom_Muscle",
            "Tomtom_Spleen"
        ],

        "source_n": 18
    },

    "S6": {
        "file":
            "Supplementary_Table_S6.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "FIMO_q005_links",
            "Motif_target_summary",
            "Tissue_network_summary",
            "High_RNA_integrated",
            "FIMO_QC"
        ],

        "source_n": 5
    },

    "S7": {
        "file":
            "Supplementary_Table_S7.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Candidate_TF_master"
        ],

        "source_n": 1
    },

    "S8": {
        "file":
            "Supplementary_Table_S8.xlsx",

        "sheets": [
            "README",
            "SOURCE_INDEX",
            "Module_summary",
            "Module_definitions",
            "Module_peak_gene_pairs",
            "Gene_ranking",
            "Top10_representatives",
            "Top5_primary_supporting",
            "Within_tissue_overlap",
            "Representative_edges",
            "Representative_nodes"
        ],

        "source_n": 9
    }
}


# =====================================================================
# Critical semantic wording
# =====================================================================

SEMANTIC_RULES = {

    "S2": [
        "FDR < 0.01",
        "linear tissue-mean TMM.TPM",
        "continuous tissue-preference",
        "frozen Step03"
    ],

    "S3": [
        "strong coupling",
        "genomic associations",
        "not be interpreted as experimentally validated causal interactions"
    ],

    "S4": [
        "-2000/+500",
        "Promoter > 5UTR > 3UTR > Exon > Intron > Distal",
        "do not establish regulatory causality"
    ],

    "S5": [
        "AME formal significance criterion was E < 0.05",
        "no motif met this formal criterion",
        "0/9 met E < 0.05",
        "training-score-only",
        "not independent replication"
    ],

    "S6": [
        "612 FIMO q <= 0.05",
        "sequence-match",
        "must not be interpreted as direct TF occupancy",
        "not the primary FIMO evidence table",
        "does not by itself establish binding"
    ],

    "S7": [
        "48 frozen tissue–TF candidate records",
        "do not demonstrate direct DNA binding",
        "Candidate TF"
    ],

    "S8": [
        "candidate regulatory modules",
        "not be interpreted as validated direct regulatory interactions",
        "do not establish TF cooperation",
        "visualization-oriented subset"
    ]
}


# =====================================================================
# Helpers
# =====================================================================

def md5(path):

    h = hashlib.md5()

    with open(path, "rb") as f:

        for block in iter(
            lambda: f.read(1024 * 1024),
            b""
        ):
            h.update(block)

    return h.hexdigest()


def read_tsv(path):

    with open(
        path,
        "r",
        encoding="utf-8",
        errors="replace"
    ) as f:

        return list(
            csv.DictReader(
                f,
                delimiter="\t"
            )
        )


def worksheet_text(ws):

    texts = []

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is not None:

                texts.append(
                    str(cell.value)
                )

    return "\n".join(texts)


# =====================================================================
# Read frozen manifests
# =====================================================================

source_lock_rows = read_tsv(
    SOURCE_LOCK
)

build_manifest_rows = read_tsv(
    BUILD_MANIFEST
)

sheet_qc_rows = read_tsv(
    BUILD_SHEET_QC
)

primary_qc_rows = read_tsv(
    BUILD_PRIMARY_QC
)

build_md5_rows = read_tsv(
    BUILD_MD5
)


if len(source_lock_rows) != 55:

    raise SystemExit(
        "ERROR: Step10F4A2 source lock "
        f"contains {len(source_lock_rows)} rows, expected 55."
    )


# =====================================================================
# 1. Verify all 55 source files remain frozen
# =====================================================================

source_trace_rows = []

source_fail = False

for x in source_lock_rows:

    p = ROOT / x["Path"]

    exists = p.exists()

    current_md5 = (
        md5(p)
        if exists
        else ""
    )

    status = (
        "PASS"
        if (
            exists
            and
            current_md5 == x["MD5"]
        )
        else "FAIL"
    )

    if status != "PASS":
        source_fail = True

    source_trace_rows.append({
        "Table":
            x["Table"],

        "Sheet":
            x["Sheet"],

        "Source_path":
            x["Path"],

        "Locked_MD5":
            x["MD5"],

        "Current_MD5":
            current_md5,

        "Status":
            status
    })


# =====================================================================
# 2. Exact workbook inventory
# =====================================================================

actual_xlsx = sorted(
    x.name
    for x in BUILD_B.glob(
        "*.xlsx"
    )
)

expected_xlsx = sorted(
    x["file"]
    for x in EXPECTED.values()
)

inventory_status = (
    "PASS"
    if actual_xlsx == expected_xlsx
    else "FAIL"
)


# =====================================================================
# 3. Frozen workbook MD5
# =====================================================================

locked_workbook_md5 = {}

for row in build_md5_rows:

    locked_workbook_md5[
        row["Table"]
    ] = row["MD5"]


# =====================================================================
# 4. Workbook technical + traceability + semantic QC
# =====================================================================

technical_rows = []
semantic_rows = []
traceability_rows = []

overall_fail = (
    source_fail
    or inventory_status != "PASS"
)


# Build expected source tuples by Table.
source_by_table = {}

for x in source_lock_rows:

    source_by_table.setdefault(
        x["Table"],
        []
    ).append(x)


# Expected data-row counts from Step10F4B.
sheet_qc_lookup = {}

for x in sheet_qc_rows:

    sheet_qc_lookup[
        (
            x["Table"],
            x["Sheet"]
        )
    ] = x


for table in [
    "S1",
    "S2",
    "S3",
    "S4",
    "S5",
    "S6",
    "S7",
    "S8"
]:

    spec = EXPECTED[
        table
    ]

    workbook_path = (
        BUILD_B /
        spec["file"]
    )

    exists = workbook_path.exists()

    current_md5 = (
        md5(workbook_path)
        if exists
        else ""
    )

    expected_md5 = locked_workbook_md5.get(
        table,
        ""
    )

    md5_status = (
        "PASS"
        if (
            exists
            and
            expected_md5
            and
            current_md5 == expected_md5
        )
        else "FAIL"
    )

    if not exists:

        overall_fail = True

        technical_rows.append({
            "Table": table,
            "Workbook": spec["file"],
            "Check": "Workbook_exists",
            "Observed": 0,
            "Expected": 1,
            "Status": "FAIL"
        })

        continue


    wb = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_links=False
    )


    # ---------------------------------------------------------
    # Sheet names / order
    # ---------------------------------------------------------

    actual_sheets = wb.sheetnames

    sheet_status = (
        "PASS"
        if actual_sheets == spec["sheets"]
        else "FAIL"
    )

    if sheet_status != "PASS":
        overall_fail = True


    # ---------------------------------------------------------
    # Hidden sheets
    # ---------------------------------------------------------

    hidden = [
        ws.title
        for ws in wb.worksheets
        if ws.sheet_state != "visible"
    ]

    hidden_status = (
        "PASS"
        if len(hidden) == 0
        else "FAIL"
    )

    if hidden_status != "PASS":
        overall_fail = True


    # ---------------------------------------------------------
    # Unexpected formulas
    # ---------------------------------------------------------

    formula_cells = 0

    for ws in wb.worksheets:

        if ws.title in {
            "README",
            "SOURCE_INDEX"
        }:
            continue

        for row in ws.iter_rows():

            for cell in row:

                if cell.data_type == "f":
                    formula_cells += 1

    formula_status = (
        "PASS"
        if formula_cells == 0
        else "FAIL"
    )

    if formula_status != "PASS":
        overall_fail = True


    # ---------------------------------------------------------
    # Data sheet row counts
    # ---------------------------------------------------------

    row_count_failures = []

    for sheet in spec["sheets"]:

        if sheet in {
            "README",
            "SOURCE_INDEX"
        }:
            continue

        if sheet not in wb.sheetnames:

            row_count_failures.append(
                f"{sheet}:MISSING"
            )
            continue

        ws = wb[
            sheet
        ]

        observed = max(
            ws.max_row - 1,
            0
        )

        q = sheet_qc_lookup.get(
            (
                table,
                sheet
            )
        )

        expected = (
            int(
                q["Workbook_data_rows"]
            )
            if q is not None
            else None
        )

        if (
            expected is None
            or
            observed != expected
        ):

            row_count_failures.append(
                f"{sheet}:{observed}/{expected}"
            )


    rows_status = (
        "PASS"
        if not row_count_failures
        else "FAIL"
    )

    if rows_status != "PASS":
        overall_fail = True


    # ---------------------------------------------------------
    # SOURCE_INDEX traceability
    # ---------------------------------------------------------

    source_index_status = "FAIL"
    source_index_n = 0
    expected_source_n = spec[
        "source_n"
    ]

    source_missing = []
    source_md5_mismatch = []


    if "SOURCE_INDEX" in wb.sheetnames:

        ws_src = wb[
            "SOURCE_INDEX"
        ]

        headers = [
            c.value
            for c in ws_src[1]
        ]

        idx = {
            str(h): i
            for i, h in enumerate(headers)
            if h is not None
        }

        needed = [
            "Sheet",
            "Source_path",
            "Locked_MD5"
        ]

        if all(
            x in idx
            for x in needed
        ):

            observed_sources = {}

            for row in ws_src.iter_rows(
                min_row=2,
                values_only=True
            ):

                path = row[
                    idx["Source_path"]
                ]

                locked = row[
                    idx["Locked_MD5"]
                ]

                sheet_role = row[
                    idx["Sheet"]
                ]

                if path is None:
                    continue

                observed_sources[
                    str(path)
                ] = {
                    "md5":
                        str(locked)
                        if locked is not None
                        else "",
                    "sheet":
                        str(sheet_role)
                }


            source_index_n = len(
                observed_sources
            )


            for expected_source in source_by_table[
                table
            ]:

                p = expected_source[
                    "Path"
                ]

                if p not in observed_sources:

                    source_missing.append(
                        p
                    )

                else:

                    if (
                        observed_sources[p]["md5"]
                        != expected_source["MD5"]
                    ):

                        source_md5_mismatch.append(
                            p
                        )


            source_index_status = (
                "PASS"
                if (
                    source_index_n
                    == expected_source_n
                    and
                    not source_missing
                    and
                    not source_md5_mismatch
                )
                else "FAIL"
            )


    if source_index_status != "PASS":
        overall_fail = True


    traceability_rows.append({
        "Table":
            table,

        "Expected_locked_sources":
            expected_source_n,

        "SOURCE_INDEX_rows":
            source_index_n,

        "Missing_source_paths":
            ";".join(
                source_missing
            ),

        "MD5_mismatches":
            ";".join(
                source_md5_mismatch
            ),

        "Status":
            source_index_status
    })


    # ---------------------------------------------------------
    # README semantic lock
    # ---------------------------------------------------------

    semantic_status = "PASS"

    missing_phrases = []

    if table in SEMANTIC_RULES:

        if "README" not in wb.sheetnames:

            semantic_status = "FAIL"

            missing_phrases.append(
                "README_MISSING"
            )

        else:

            txt = worksheet_text(
                wb["README"]
            ).lower()

            for phrase in SEMANTIC_RULES[
                table
            ]:

                if phrase.lower() not in txt:

                    missing_phrases.append(
                        phrase
                    )


            if missing_phrases:

                semantic_status = "FAIL"


    if semantic_status != "PASS":
        overall_fail = True


    semantic_rows.append({
        "Table":
            table,

        "Required_phrases":
            len(
                SEMANTIC_RULES.get(
                    table,
                    []
                )
            ),

        "Missing_phrases":
            "; ".join(
                missing_phrases
            ),

        "Status":
            semantic_status
    })


    # ---------------------------------------------------------
    # Technical manifest
    # ---------------------------------------------------------

    checks = [

        (
            "Workbook_exists",
            1,
            1,
            "PASS"
        ),

        (
            "Workbook_MD5",
            current_md5,
            expected_md5,
            md5_status
        ),

        (
            "Sheet_count",
            len(actual_sheets),
            len(spec["sheets"]),
            (
                "PASS"
                if len(actual_sheets)
                == len(spec["sheets"])
                else "FAIL"
            )
        ),

        (
            "Sheet_names_and_order",
            "|".join(
                actual_sheets
            ),
            "|".join(
                spec["sheets"]
            ),
            sheet_status
        ),

        (
            "Hidden_sheet_count",
            len(hidden),
            0,
            hidden_status
        ),

        (
            "Unexpected_formula_cells",
            formula_cells,
            0,
            formula_status
        ),

        (
            "Data_sheet_row_counts",
            (
                "PASS"
                if not row_count_failures
                else ";".join(
                    row_count_failures
                )
            ),
            "PASS",
            rows_status
        ),

        (
            "SOURCE_INDEX_traceability",
            source_index_n,
            expected_source_n,
            source_index_status
        )
    ]


    for (
        check,
        observed,
        expected,
        status
    ) in checks:

        if status != "PASS":
            overall_fail = True

        technical_rows.append({
            "Table":
                table,

            "Workbook":
                spec["file"],

            "Check":
                check,

            "Observed":
                observed,

            "Expected":
                expected,

            "Status":
                status
        })


    wb.close()


# =====================================================================
# 5. Re-check Step10F4B primary counts
# =====================================================================

primary_rows = []

for x in primary_qc_rows:

    status = (
        "PASS"
        if (
            x["Status"] == "PASS"
            and
            x["Observed"] == x["Expected"]
        )
        else "FAIL"
    )

    if status != "PASS":
        overall_fail = True

    primary_rows.append({
        "Metric":
            x["Metric"],

        "Observed":
            x["Observed"],

        "Expected":
            x["Expected"],

        "Status":
            status
    })


# =====================================================================
# 6. Write manifests
# =====================================================================

technical_file = (
    OUTDIR /
    "10F4C_technical_lock.tsv"
)

with open(
    technical_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Table",
            "Workbook",
            "Check",
            "Observed",
            "Expected",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        technical_rows
    )


trace_file = (
    OUTDIR /
    "10F4C_source_traceability.tsv"
)

with open(
    trace_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Table",
            "Expected_locked_sources",
            "SOURCE_INDEX_rows",
            "Missing_source_paths",
            "MD5_mismatches",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        traceability_rows
    )


source_file = (
    OUTDIR /
    "10F4C_55_source_MD5_recheck.tsv"
)

with open(
    source_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Table",
            "Sheet",
            "Source_path",
            "Locked_MD5",
            "Current_MD5",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        source_trace_rows
    )


semantic_file = (
    OUTDIR /
    "10F4C_semantic_lock.tsv"
)

with open(
    semantic_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Table",
            "Required_phrases",
            "Missing_phrases",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        semantic_rows
    )


primary_file = (
    OUTDIR /
    "10F4C_primary_count_lock.tsv"
)

with open(
    primary_file,
    "w",
    newline=""
) as f:

    writer = csv.DictWriter(
        f,
        fieldnames=[
            "Metric",
            "Observed",
            "Expected",
            "Status"
        ],
        delimiter="\t"
    )

    writer.writeheader()
    writer.writerows(
        primary_rows
    )


# =====================================================================
# 7. Final MD5 lock manifest
# =====================================================================

final_md5_file = (
    OUTDIR /
    "10F4C_FINAL_workbook_MD5.tsv"
)

with open(
    final_md5_file,
    "w"
) as f:

    f.write(
        "Table\tWorkbook\tMD5\n"
    )

    for table in [
        "S1",
        "S2",
        "S3",
        "S4",
        "S5",
        "S6",
        "S7",
        "S8"
    ]:

        p = (
            BUILD_B /
            EXPECTED[table]["file"]
        )

        f.write(
            f"{table}\t"
            f"{p}\t"
            f"{md5(p) if p.exists() else 'MISSING'}\n"
        )


# =====================================================================
# 8. Overall status
# =====================================================================

overall = (
    "FAIL"
    if overall_fail
    else "PASS"
)


status_file = (
    OUTDIR /
    "10F4C_overall_status.tsv"
)

with open(
    status_file,
    "w"
) as f:

    f.write(
        "Metric\tValue\n"
    )

    f.write(
        f"Workbook_inventory\t"
        f"{inventory_status}\n"
    )

    f.write(
        "Locked_source_files_expected\t55\n"
    )

    f.write(
        "Locked_source_files_MD5_PASS\t"
        f"{sum(x['Status'] == 'PASS' for x in source_trace_rows)}\n"
    )

    f.write(
        "Final_workbooks_expected\t8\n"
    )

    f.write(
        "Final_workbooks_present\t"
        f"{len(actual_xlsx)}\n"
    )

    f.write(
        f"Supplementary_Tables_S1_S8_consistency_lock\t"
        f"{overall}\n"
    )

    f.write(
        "Biological_results_modified\t0\n"
    )


# =====================================================================
# 9. LOCK marker only after complete PASS
# =====================================================================

lock_marker = (
    OUTDIR /
    "SUPPLEMENTARY_TABLES_S1_S8_LOCKED.txt"
)

if overall == "PASS":

    with open(
        lock_marker,
        "w"
    ) as f:

        f.write(
            "Supplementary Tables S1-S8 are FINAL and LOCKED.\n"
        )

        f.write(
            "Step10F4C consistency lock: PASS\n"
        )

        f.write(
            "Frozen source files verified: 55/55\n"
        )

        f.write(
            "Final workbooks verified: 8/8\n"
        )

        f.write(
            "Biological results modified: 0\n"
        )

else:

    if lock_marker.exists():
        lock_marker.unlink()


# =====================================================================
# 10. Console report
# =====================================================================

print()
print("=" * 92)
print("WORKBOOK TECHNICAL LOCK SUMMARY")
print("=" * 92)

for table in [
    "S1",
    "S2",
    "S3",
    "S4",
    "S5",
    "S6",
    "S7",
    "S8"
]:

    rows = [
        x for x in technical_rows
        if x["Table"] == table
    ]

    p = sum(
        x["Status"] == "PASS"
        for x in rows
    )

    print(
        f"{table}: "
        f"{p}/{len(rows)} technical checks PASS"
    )


print()
print("=" * 92)
print("SOURCE TRACEABILITY")
print("=" * 92)

for x in traceability_rows:

    print(
        f"{x['Table']}: "
        f"{x['SOURCE_INDEX_rows']}/"
        f"{x['Expected_locked_sources']} sources "
        f"{x['Status']}"
    )


print()
print("=" * 92)
print("SEMANTIC LOCK")
print("=" * 92)

for x in semantic_rows:

    print(
        f"{x['Table']}: "
        f"{x['Status']}"
    )

    if x["Missing_phrases"]:

        print(
            "    Missing:",
            x["Missing_phrases"]
        )


print()
print("=" * 92)
print("PRIMARY COUNT LOCK")
print("=" * 92)

for x in primary_rows:

    print(
        f"{x['Metric']:<43} "
        f"{str(x['Observed']):>6} / "
        f"{str(x['Expected']):<6} "
        f"{x['Status']}"
    )


print()
print("=" * 92)
print(
    f"STEP10F4C STATUS: {overall}"
)

if overall == "PASS":

    print(
        "Supplementary Tables S1–S8 are technically, "
        "semantically and source-traceably LOCKED."
    )

    print(
        "Do not modify the final workbooks after this point."
    )

else:

    print(
        "Supplementary Tables are NOT locked."
    )

    print(
        "Resolve failed checks before proceeding."
    )

print(
    "Biological results modified: 0"
)

print("=" * 92)

print("\nOutputs:")
print(technical_file)
print(trace_file)
print(source_file)
print(semantic_file)
print(primary_file)
print(final_md5_file)
print(status_file)

if overall == "PASS":
    print(lock_marker)

